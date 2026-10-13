"""A cross-tabulation in an export that writes tables.

A crosstab's cube is one row per combination of everything the page can
filter on, which is how the exported page re-adds it as a reader narrows the
board. It is not how anybody reads a cross-tabulation, and a report that
wrote the cube out as a list would be a table nobody could quote from. These
cover laying the grid back out: the cells added up, the percentages taken
against the right margin, the totals where the widget asks for them, and a
withheld figure left withheld with the mark explained under the table.

The export wrote the grid out as nothing at all until now - "it carries no
numbers", under the heading of a table full of them - which is the failure
these exist to keep from coming back.
"""

from __future__ import annotations

import io

import pandas as pd
import pytest

from app.services import export_document
from app.services.export_tables import crosstab_table, disclosure_note, widget_table
from tests.test_api_analytics import _stata_bytes, _zip_bytes

CUBE = {
    "columns": [
        {"name": "province", "label": "Province"},
        {"name": "sex", "label": "Sex"},
        {"name": "__value", "label": "Interviews"},
    ],
    "dimensions": ["province", "sex"],
    "measures": [
        {"name": "__value", "label": "Interviews", "how": "sum", "parts": ["__value"]}
    ],
    # Out of order on purpose: a cube arrives in whatever order the query
    # answered in, and a table is read down its own labels.
    "rows": [
        ["Southern", "Male", 6],
        ["Northern", "Male", 12],
        ["Southern", "Female", 4],
        ["Northern", "Female", 8],
    ],
    "sort": [],
}


def _widget(cube=None, **crosstab) -> dict:
    spec = {
        "row_variable": "Province",
        "column_variable": "Sex",
        "percentages": "none",
        "include_totals": True,
        "show_values": True,
        "measure_label": "Count",
    }
    spec.update(crosstab)
    return {
        "title": "Sex by province",
        "type": "crosstab",
        "kind": "crosstab",
        "page": 0,
        "cube": CUBE if cube is None else cube,
        "crosstab": spec,
    }


def test_it_comes_out_as_a_grid_rather_than_a_list_of_cells():
    heading, rows = crosstab_table(_widget())

    assert heading == ["Province", "Female", "Male", "Total"]
    assert rows == [
        ["Northern", 8.0, 12.0, 20.0],
        ["Southern", 4.0, 6.0, 10.0],
        ["Total", 12.0, 18.0, 30.0],
    ]


def test_a_crosstab_widget_is_not_reported_as_having_no_numbers():
    """What the Word report did with every crosstab on a board until now."""
    assert widget_table(_widget()) is not None


def test_two_kinds_of_nothing_are_one_row_of_missing():
    """A blank and a null are both "(missing)", so the cell has to add them.

    Two rows of the cube land in the same cell of the grid, and a cell that
    took the second rather than both would report fewer missing answers than
    the dataset holds - in the one column somebody is reading to find out how
    many answers are missing.
    """
    cube = dict(CUBE, rows=[
        [None, "Male", 7],
        ["", "Male", 5],
        ["Northern", "Male", 12],
        ["Northern", "Female", 8],
    ])

    _, rows = crosstab_table(_widget(cube=cube, include_totals=False))

    assert rows[0] == ["(missing)", 0.0, 12.0]


@pytest.mark.parametrize(
    ("percentages", "northern_female"),
    [
        # 8 of Northern's 20, 8 of the 12 women, 8 of all 30.
        ("row", "40.0%"),
        ("column", "66.7%"),
        ("total", "26.7%"),
    ],
)
def test_a_percentage_is_taken_against_the_margin_it_names(percentages, northern_female):
    _, rows = crosstab_table(_widget(percentages=percentages, include_totals=False))

    assert rows[0][:2] == ["Northern", northern_female]


def test_a_percentage_of_nothing_is_said_rather_than_divided():
    cube = dict(CUBE, rows=[["Northern", "Male", 0], ["Northern", "Female", 0]])

    _, rows = crosstab_table(_widget(cube=cube, percentages="row", include_totals=False))

    assert rows[0] == ["Northern", "-", "-"]


def test_totals_are_only_there_where_the_widget_shows_them():
    heading, rows = crosstab_table(_widget(include_totals=False))

    assert heading == ["Province", "Female", "Male"]
    assert [row[0] for row in rows] == ["Northern", "Southern"]


def test_a_table_of_categories_alone_has_no_cells_to_write():
    """A one-way table asked for with no measure: the categories, nothing else."""
    heading, rows = crosstab_table(_widget(show_values=False, measure_label=""))

    assert heading == ["Province"]
    assert rows == [["Northern"], ["Southern"]]


def test_a_one_way_table_heads_its_one_column_with_the_measure():
    """No column variable: the categories down the side, their counts beside."""
    cube = {
        "columns": [
            {"name": "province", "label": "Province"},
            {"name": "__value", "label": "Interviews"},
        ],
        "dimensions": ["province"],
        "measures": [
            {"name": "__value", "label": "Interviews", "how": "sum", "parts": ["__value"]}
        ],
        "rows": [["Northern", 20], ["Southern", 10]],
        "sort": [],
    }

    heading, rows = crosstab_table(_widget(cube=cube, include_totals=False))

    assert heading == ["Province", "Count"]
    assert rows == [["Northern", 20.0], ["Southern", 10.0]]


def test_a_day_is_a_day_rather_than_a_moment_at_midnight():
    """A date grouping comes back as a timestamp; a row heading is read."""
    cube = {
        "columns": [
            {"name": "day", "label": "Day"},
            {"name": "__value", "label": "Interviews"},
        ],
        "dimensions": ["day"],
        "measures": [
            {"name": "__value", "label": "Interviews", "how": "sum", "parts": ["__value"]}
        ],
        "rows": [["2026-07-19 00:00:00", 4]],
        "sort": [],
    }

    _, rows = crosstab_table(_widget(cube=cube, include_totals=False))

    assert rows[0][0] == "2026-07-19"


# --- Under disclosure control ------------------------------------------------
#
# Where a floor is set the table is computed on the server and travels
# finished, because a cube is the cell values and the withheld ones are
# exactly what must not travel. So nothing here is recomputed either: a
# withheld cell stays a mark, and the mark gets explained.


FIXED = {
    "row_variable": "Province",
    "column_variable": "Sex",
    "percentages": "none",
    "include_totals": True,
    "show_values": True,
    "measure_label": "",
    "fixed": {
        "row_labels": ["Northern", "Southern"],
        "column_labels": ["Male", "Female"],
        "values": [[12, 8], [3, None]],
        "suppressed": [[False, False], [True, False]],
        "row_totals": [20, None],
        "column_totals": [15, 8],
        "grand_total": 23,
        "row_totals_suppressed": [False, True],
        "column_totals_suppressed": [False, False],
        "grand_total_suppressed": False,
        "threshold": 5,
        "rows_withheld": 2,
    },
}


def test_a_withheld_cell_is_not_written_out():
    heading, rows = crosstab_table({"kind": "crosstab", "cube": {}, "crosstab": FIXED})

    assert heading == ["Province", "Male", "Female", "Total"]
    assert rows[0] == ["Northern", 12, 8, 20]
    # The 3 behind the mark is the number the floor exists to keep back, and
    # a report is the copy nobody can withdraw.
    assert rows[1] == ["Southern", "*", "-", "*"]
    assert rows[2] == ["Total", 15, 8, 23]
    assert "3" not in str(rows[1])


def test_a_table_of_one_column_is_not_given_a_total_of_itself():
    """One column crossed with nothing: its own total beside it says nothing."""
    spec = {
        **FIXED,
        "fixed": {
            **FIXED["fixed"],
            "column_labels": ["Male"],
            "values": [[12], [3]],
            "suppressed": [[False], [False]],
            "column_totals": [15],
        },
    }

    heading, rows = crosstab_table({"kind": "crosstab", "cube": {}, "crosstab": spec})

    assert heading == ["Province", "Male"]
    assert rows[0] == ["Northern", 12]


def test_a_withheld_table_of_percentages_keeps_the_sign_on_each_cell():
    spec = {
        **FIXED,
        "percentages": "row",
        "fixed": {
            **FIXED["fixed"],
            "values": [[60, 40], [75, None]],
            "row_totals": [100, None],
            "column_totals": [67, 33],
            "grand_total": 100,
        },
    }

    _, rows = crosstab_table({"kind": "crosstab", "cube": {}, "crosstab": spec})

    assert rows[0] == ["Northern", "60%", "40%", 100]
    assert rows[1] == ["Southern", "*", "-", "*"]


def test_the_mark_is_explained_under_the_table():
    note = disclosure_note({"crosstab": FIXED})

    assert "* withheld to protect fewer than 5 records." in note
    assert "2 categories below that floor are not listed." in note


def test_one_category_is_said_in_the_singular():
    spec = {**FIXED, "fixed": {**FIXED["fixed"], "rows_withheld": 1}}

    assert "1 category below that floor is not listed." in disclosure_note({"crosstab": spec})


def test_a_table_with_nothing_withheld_says_nothing():
    spec = {
        **FIXED,
        "fixed": {
            **FIXED["fixed"],
            "values": [[12, 8], [3, 4]],
            "suppressed": [[False, False], [False, False]],
            "rows_withheld": 0,
        },
    }

    assert disclosure_note({"crosstab": spec}) == ""
    assert disclosure_note(_widget()) == ""


# --- Through the document ----------------------------------------------------


def _read(content: bytes):
    from docx import Document

    return Document(io.BytesIO(content))


def test_the_document_carries_the_grid_and_its_note():
    payload = {
        "name": "Census board",
        "generated_at": "2026-10-12T09:30:00+00:00",
        "pages": [{"name": "Home"}],
        "widgets": [{**_widget(cube={}, **FIXED), "title": "Sex by province"}],
    }

    document = _read(export_document.document(payload))

    table = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
    assert table[0] == ["Province", "Male", "Female", "Total"]
    assert table[2] == ["Southern", "*", "-", "*"]
    assert any("withheld to protect fewer than 5" in p.text for p in document.paragraphs)


def test_a_frozen_widget_says_why_rather_than_that_it_has_no_numbers():
    """A chart on a median has numbers; they cannot be worked out again here."""
    payload = {
        "name": "Board",
        "generated_at": "2026-10-12T09:30:00+00:00",
        "pages": [{"name": "Home"}],
        "widgets": [
            {
                "title": "Median age",
                "type": "chart",
                "kind": "snapshot",
                "page": 0,
                "frozen_because": "a median cannot be worked out again from group totals",
            },
            {"title": "Where the work is", "type": "map", "kind": "snapshot", "page": 0},
        ],
    }

    text = "\n".join(p.text for p in _read(export_document.document(payload)).paragraphs)

    assert "a median cannot be worked out again from group totals" in text
    assert "a map of locations" in text


def test_a_board_with_a_crosstab_reports_the_same_table_the_data_holds(
    client, auth_headers
):
    """End to end, against the figures pandas gets from the same rows."""
    frame = pd.DataFrame(
        {
            "interview__key": [f"k{i}" for i in range(12)],
            "province": ["Northern"] * 7 + ["Highland"] * 5,
            "sex": ["Male", "Female"] * 6,
        }
    )
    uploaded = client.post(
        "/api/v1/datasets/upload",
        headers=auth_headers,
        files={"file": ("tab.zip", _zip_bytes({"tab.dta": _stata_bytes(frame)}),
                        "application/zip")},
    )
    dataset_id = uploaded.json()["datasets"][0]["id"]
    board = client.post(
        "/api/v1/dashboards", headers=auth_headers, json={"name": "Tabulations"}
    ).json()
    chart = client.post(
        "/api/v1/dashboards/charts",
        headers=auth_headers,
        json={
            "name": "Sex by province",
            "dataset_id": dataset_id,
            "chart_type": "crosstab",
            "spec": {"crosstab": {"row_variable": "province", "column_variable": "sex"}},
        },
    ).json()
    client.post(
        f"/api/v1/dashboards/{board['id']}/widgets",
        headers=auth_headers,
        json={"title": "Sex by province", "widget_type": "chart", "chart_id": chart["id"]},
    )

    downloaded = client.get(
        f"/api/v1/dashboards/{board['id']}/export.docx", headers=auth_headers
    )
    assert downloaded.status_code == 200, downloaded.text

    rows = [
        [cell.text for cell in row.cells] for row in _read(downloaded.content).tables[0].rows
    ]
    expected = pd.crosstab(frame.province, frame.sex)
    written = {
        row[0]: dict(zip(rows[0][1:], row[1:], strict=True))
        for row in rows[1:]
        if row[0] != "Total"
    }
    for province in expected.index:
        for sex in expected.columns:
            assert written[province][sex] == str(expected.loc[province, sex])
        assert written[province]["Total"] == str(expected.loc[province].sum())
