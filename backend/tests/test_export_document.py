"""A board as a Word report.

The web page is the board and the workbook is its figures. This is the one
that goes inside a written report, so what matters is that it reads as a
document: the board's name and what it was narrowed to, then each page, then
each widget as a heading and a table somebody can quote from.
"""

from __future__ import annotations

import io

import pandas as pd
import pytest

from app.services import export_document
from tests.test_api_analytics import _stata_bytes, _zip_bytes


def _read(content: bytes):
    from docx import Document

    return Document(io.BytesIO(content))


def _text(document) -> str:
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def _payload(**extra):
    return {
        "name": "Cook Islands Census 2026",
        "description": "Fieldwork monitoring",
        "generated_at": "2026-10-12T09:30:00+00:00",
        "pages": [{"name": "Home Page"}],
        "widgets": [],
        **extra,
    }


def test_it_opens_as_a_document_and_says_what_it_is_of():
    document = _read(export_document.document(_payload(
        pinned=[{"variable": "island", "label": "Island", "value": "Pukapuka"}],
    )))
    text = _text(document)

    assert "Cook Islands Census 2026" in text
    assert "Fieldwork monitoring" in text
    assert "Island: Pukapuka" in text
    assert "12 October 2026, 09:30 UTC" in text
    assert "as they stood then" in text


def test_a_widget_becomes_a_heading_and_a_table():
    cube = {
        "dimensions": ["province"],
        "filters": [],
        "measures": [{"name": "value", "label": "Interviews", "how": "sum",
                      "parts": ["value"]}],
        "columns": [
            {"name": "province", "label": "Province", "type": "dimension", "data_type": "text"},
            {"name": "value", "label": "Interviews", "type": "measure", "data_type": "number"},
        ],
        "rows": [["Shefa", 1240], ["Sanma", 860]],
        "top": None,
        "sort": [],
    }
    document = _read(export_document.document(_payload(widgets=[
        {"title": "Interviews by province", "type": "chart", "kind": "chart",
         "page": 0, "cube": cube},
    ])))

    assert "Interviews by province" in [p.text for p in document.paragraphs]
    assert len(document.tables) == 1
    rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
    assert rows[0] == ["Province", "Interviews"]
    # Thousands separated, and no decimal point on a count.
    assert rows[1] == ["Shefa", "1,240"]
    assert rows[2] == ["Sanma", "860"]


def test_a_text_widget_is_carried_as_what_it_says():
    """It is already prose; a document is where prose belongs."""
    document = _read(export_document.document(_payload(widgets=[
        {"title": "About this board", "type": "text", "kind": "snapshot", "page": 0,
         "config": {"content": "Updated every morning from the tablets."}},
    ])))
    text = _text(document)

    assert "Updated every morning from the tablets." in text
    assert not document.tables


def test_a_widget_with_no_numbers_says_so_rather_than_vanishing():
    document = _read(export_document.document(_payload(widgets=[
        {"title": "Where the work is", "type": "map", "kind": "snapshot", "page": 0},
    ])))
    text = _text(document)

    assert "Where the work is" in text
    assert "a map of locations" in text


def test_pages_are_headed_only_when_there_is_more_than_one():
    """A board of one page is already called that by the document's title."""
    one = _read(export_document.document(_payload(widgets=[
        {"title": "A", "type": "text", "kind": "snapshot", "page": 0,
         "config": {"content": "first"}},
    ])))
    assert "Home Page" not in _text(one)

    two = _read(export_document.document(_payload(
        pages=[{"name": "Home Page"}, {"name": "Fieldwork"}],
        widgets=[
            {"title": "A", "type": "text", "kind": "snapshot", "page": 0,
             "config": {"content": "first"}},
            {"title": "B", "type": "text", "kind": "snapshot", "page": 1,
             "config": {"content": "second"}},
        ],
    )))
    headings = _text(two)
    assert "Home Page" in headings and "Fieldwork" in headings


def test_the_document_and_the_board_hold_the_same_figures(client, auth_headers):
    frame = pd.DataFrame(
        {
            "interview__key": [f"k{i}" for i in range(10)],
            "province": ["Shefa"] * 6 + ["Sanma"] * 4,
            "age": [20.0, 30.0, 40.0, 50.0, 60.0, 10.0, 25.0, 35.0, 45.0, 55.0],
        }
    )
    uploaded = client.post(
        "/api/v1/datasets/upload",
        headers=auth_headers,
        files={"file": ("report.zip", _zip_bytes({"report.dta": _stata_bytes(frame)}),
                        "application/zip")},
    )
    dataset_id = uploaded.json()["datasets"][0]["id"]
    board = client.post(
        "/api/v1/dashboards", headers=auth_headers, json={"name": "Report board"}
    ).json()
    chart = client.post(
        "/api/v1/dashboards/charts",
        headers=auth_headers,
        json={
            "name": "Mean age by province",
            "dataset_id": dataset_id,
            "chart_type": "bar",
            "spec": {"query": {
                "dimensions": [{"variable": "province"}],
                "measures": [{"agg": "mean", "variable": "age"}],
            }},
        },
    ).json()
    client.post(
        f"/api/v1/dashboards/{board['id']}/widgets",
        headers=auth_headers,
        json={"title": "Mean age by province", "widget_type": "chart",
              "chart_id": chart["id"]},
    )

    downloaded = client.get(
        f"/api/v1/dashboards/{board['id']}/export.docx", headers=auth_headers
    )
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.headers["content-disposition"].endswith('report-board.docx"')

    rows = [
        [cell.text for cell in row.cells] for row in _read(downloaded.content).tables[0].rows
    ]
    # Read back as numbers: how they are written is the unit test above, and
    # what matters here is that the document holds the board's own figures.
    means = {row[0]: float(row[1].replace(",", "")) for row in rows[1:]}
    assert means["Shefa"] == pytest.approx(frame[frame.province == "Shefa"]["age"].mean())
    assert means["Sanma"] == pytest.approx(frame[frame.province == "Sanma"]["age"].mean())


def test_a_number_with_a_decimal_keeps_it():
    assert export_document._written(1240) == "1,240"
    assert export_document._written(1240.0) == "1,240"
    assert export_document._written(36.6666) == "36.67"
    assert export_document._written(None) == ""
    assert export_document._written(float("nan")) == ""
    assert export_document._written("Shefa") == "Shefa"


@pytest.mark.parametrize("value", [True, False])
def test_a_yes_is_not_a_one(value):
    assert export_document._written(value) == ("Yes" if value else "No")
