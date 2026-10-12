"""A board as a document, for the report it ends up in.

The web page is the board. This is the other thing people ask an export
for: something to put in a report, or to send to somebody who will read it
rather than work with it. A Word file, because that is what a ministry's
report is written in and a PDF cannot be edited into one.

What it carries is the board read out in order: the title and what it was
narrowed to, then each page, then each widget as a heading and a table of its
numbers. A chart is its figures here rather than a picture of it - a picture
of a chart in a document nobody can check is worth less than the numbers
somebody can - and a text widget is carried as what it says, because that is
already prose in a document.
"""

from __future__ import annotations

import io
from typing import Any

from app.services.export_tables import NO_NUMBERS, _said_plainly, widget_table


def _written(value: Any) -> str:
    """A number as a document writes one.

    Thousands separated, and decimals only where there are any: a column of
    counts with ".0" after each one reads as a measurement rather than a
    count.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        number = float(value)
        if number != number:  # NaN
            return ""
        if number == int(number) and abs(number) < 1e15:
            return f"{int(number):,}"
        return f"{number:,.2f}"
    return str(value)


def document(payload: dict[str, Any]) -> bytes:
    """The board as a .docx report."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    board = Document()
    muted = RGBColor(0x6B, 0x72, 0x80)

    board.add_heading(str(payload.get("name") or "Dashboard"), level=0)
    if payload.get("description"):
        board.add_paragraph(str(payload["description"]))
    for pin in payload.get("pinned") or []:
        line = board.add_paragraph()
        run = line.add_run(
            f"{pin.get('label') or pin.get('variable')}: {pin.get('value')}"
        )
        run.bold = True

    stamp = board.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = stamp.add_run(
        "Exported from SurveyHQ on "
        + _said_plainly(str(payload.get("generated_at") or ""))
        + ". The numbers are as they stood then."
    )
    run.font.size = Pt(9)
    run.font.color.rgb = muted

    pages = [str((page or {}).get("name") or "") for page in payload.get("pages") or []]
    widgets = payload.get("widgets") or []
    many = len([name for name in pages if name]) > 1

    for index, name in enumerate(pages or [""]):
        mine = [one for one in widgets if int(one.get("page") or 0) == index]
        if not mine:
            continue
        # A board of one page is not introduced by its own name: the document
        # is already called that.
        if many:
            board.add_heading(name or f"Page {index + 1}", level=1)
        for widget in mine:
            _widget(board, widget, muted)

    buffer = io.BytesIO()
    board.save(buffer)
    return buffer.getvalue()


def _widget(board: Any, widget: dict[str, Any], muted: Any) -> None:
    """One widget: its title, then whatever of it a document can carry."""
    from docx.shared import Pt

    title = str(widget.get("title") or "").strip()
    kind = str(widget.get("type") or "")
    if title:
        board.add_heading(title, level=2)

    if kind == "text":
        content = str((widget.get("config") or {}).get("content") or "").strip()
        if content:
            board.add_paragraph(content)
            return

    table = widget_table(widget)
    if table is None:
        line = board.add_paragraph()
        run = line.add_run(
            "Not in this document: " + NO_NUMBERS.get(kind, "it carries no numbers") + "."
        )
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = muted
        return

    heading, rows = table
    written = board.add_table(rows=1, cols=len(heading))
    written.style = "Light Grid Accent 1"
    for column, label in enumerate(heading):
        cell = written.rows[0].cells[column]
        cell.text = str(label)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = written.add_row().cells
        for column, value in enumerate(row[: len(heading)]):
            cells[column].text = _written(value)

    caption = str((widget.get("style") or {}).get("caption") or "").strip()
    if caption:
        line = board.add_paragraph()
        run = line.add_run(caption)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = muted
