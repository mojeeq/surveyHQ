#!/usr/bin/env python3
"""Build the susoDash user manual as a Word document.

The screenshots in docs/manual/ are captured from a running instance carrying
a demonstration survey, so every picture in the manual is of the real thing
rather than a mock-up. Re-run this after a release that changes the interface;
the screenshots are refreshed by scripts/capture_manual.mjs.

python-docx and Pillow are needed to build the manual and are deliberately not
in requirements.txt: the platform does not need them to run, and a deployment
should not carry the tools that wrote its documentation.

    pip install python-docx pillow
    python scripts/build_manual.py
"""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # pragma: no cover - a build-time dependency
    sys.exit("This needs python-docx and pillow: pip install python-docx pillow")

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "manual"
OUT = ROOT / "docs" / "susoDash-user-manual.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
BRAND = RGBColor(0x1D, 0x4E, 0xD8)

# The page is 6.5 inches of text between one-inch margins; a full-window
# screenshot is captured at 1440 wide, so this is the width everything gets
# unless it is a detail shot that would be silly blown up that far.
FULL = 6.3
figures: list[tuple[str, str]] = []


def style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, colour, before in (
        ("Heading 1", 20, BRAND, 22),
        ("Heading 2", 14, INK, 16),
        ("Heading 3", 11.5, INK, 12),
    ):
        heading = document.styles[name]
        heading.font.name = "Calibri"
        heading.font.size = Pt(size)
        heading.font.color.rgb = colour
        heading.font.bold = True
        heading.paragraph_format.space_before = Pt(before)
        heading.paragraph_format.space_after = Pt(5)
        heading.paragraph_format.keep_with_next = True


def rich(paragraph, text: str) -> None:
    """Write text, taking **bold** and `code` as the markers they look like."""
    import re

    for piece in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(piece)


def para(document: Document, text: str = "", style_name: str | None = None):
    paragraph = document.add_paragraph(style=style_name)
    rich(paragraph, text)
    return paragraph


def bullets(document: Document, items: list[str]) -> None:
    for item in items:
        para(document, item, "List Bullet")


def numbered(document: Document, items: list[str]) -> None:
    for item in items:
        para(document, item, "List Number")


def note(document: Document, text: str) -> None:
    """An aside, indented and greyed so it reads as an aside."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    rich(paragraph, text)
    for run in paragraph.runs:
        run.font.size = Pt(9.5)
        run.font.color.rgb = MUTED
    return paragraph


def figure(document: Document, name: str, caption: str, width: float = FULL) -> None:
    path = SHOTS / f"{name}.png"
    if not path.exists():
        print(f"  ! missing screenshot {name}")
        return
    from PIL import Image

    with Image.open(path) as image:
        ratio = image.height / image.width
    # A picture taller than the text column is scaled to fit the page instead,
    # so a full-page dashboard capture does not run off the bottom.
    if width * ratio > 8.4:
        width = 8.4 / ratio
    holder = document.add_paragraph()
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    holder.paragraph_format.space_before = Pt(8)
    holder.paragraph_format.space_after = Pt(2)
    holder.add_run().add_picture(str(path), width=Inches(width))

    figures.append((f"Figure {len(figures) + 1}", caption))
    line = document.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run(f"Figure {len(figures)}. {caption}")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    run.italic = True
    line.paragraph_format.space_after = Pt(12)


def table(document: Document, headings: list[str], rows: list[list[str]]) -> None:
    grid = document.add_table(rows=1, cols=len(headings))
    grid.style = "Light Grid Accent 1"
    for cell, heading in zip(grid.rows[0].cells, headings, strict=False):
        cell.text = ""
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
        run.font.size = Pt(9.5)
    for row in rows:
        cells = grid.add_row().cells
        for cell, value in zip(cells, row, strict=False):
            cell.text = ""
            rich(cell.paragraphs[0], value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def footer(document: Document) -> None:
    """Page numbers, because a manual without them cannot be referred to."""
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("susoDash user manual    ")
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        field = paragraph.add_run()
        field.font.size = Pt(8)
        field.font.color.rgb = MUTED
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = "PAGE"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        field._r.append(begin)
        field._r.append(instruction)
        field._r.append(end)


def cover(document: Document) -> None:
    for _ in range(6):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("susoDash")
    run.font.size = Pt(46)
    run.font.bold = True
    run.font.color.rgb = BRAND

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("User manual")
    run.font.size = Pt(20)
    run.font.color.rgb = INK

    blurb = document.add_paragraph()
    blurb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = blurb.add_run(
        "Monitoring survey data collection: importing, tabulating,\n"
        "dashboards, indicators, data quality and maps"
    )
    run.font.size = Pt(11.5)
    run.font.color.rgb = MUTED

    for _ in range(10):
        document.add_paragraph()
    stamp = document.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = stamp.add_run(dt.date.today().strftime("%d %B %Y"))
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED
    page_break(document)


def contents(document: Document) -> None:
    document.add_heading("Contents", level=1)
    para(
        document,
        "Every picture in this manual is a screenshot of a running instance "
        "carrying a demonstration labour force survey, so what you see here is "
        "what the platform actually looks like.",
    )
    for number, name in [
        ("1", "What susoDash is"),
        ("2", "Signing in and finding your way around"),
        ("3", "Projects: who can see what"),
        ("4", "Getting data in"),
        ("5", "Looking at a dataset"),
        ("6", "Explore: tabulations and charts"),
        ("7", "Dashboards"),
        ("8", "Maps and boundaries"),
        ("9", "Monitoring: indicators and targets"),
        ("10", "Data quality"),
        ("11", "Alerts"),
        ("12", "Sharing a dashboard"),
        ("13", "Taking your work away"),
        ("14", "Administration"),
        ("15", "Reference tables"),
        ("16", "If something goes wrong"),
    ]:
        line = document.add_paragraph()
        line.paragraph_format.space_after = Pt(2)
        run = line.add_run(f"{number}.  {name}")
        run.font.size = Pt(10.5)
    page_break(document)


def chapter_what(document: Document) -> None:
    document.add_heading("1. What susoDash is", level=1)
    para(
        document,
        "susoDash is a platform for watching a survey while it is still in the "
        "field. It takes the data your interviewers are collecting, and turns it "
        "into the things a survey manager needs during fieldwork: how many "
        "interviews are done and where, whether the numbers coming back look "
        "right, and which teams need a telephone call today.",
    )
    para(document, "It does four things.", "Normal")
    bullets(
        document,
        [
            "**Brings data in.** Either by connecting to a Survey Solutions "
            "server and importing on a schedule, or by uploading a Stata, SPSS, "
            "CSV or Excel file by hand.",
            "**Lets you tabulate it.** Frequencies, cross-tabulations, means and "
            "shares, with filters, weights and charts, without writing code.",
            "**Puts it on a dashboard.** Charts, tables, indicator tiles, maps "
            "and data-quality panels, arranged how you like and shared by link.",
            "**Watches it for you.** Indicators against targets, data-quality "
            "checks that run on a schedule, and alerts when something crosses a "
            "threshold you set.",
        ],
    )
    para(
        document,
        "It is self-hosted. It runs on your own server, the data never leaves "
        "your network, and it works on a machine with no route to the internet "
        "apart from the map tiles.",
    )

    document.add_heading("Who it is for", level=2)
    table(
        document,
        ["If you are", "You will mostly use"],
        [
            ["A survey manager", "Dashboards, monitoring and alerts"],
            ["A data manager or statistician", "Explore, data quality, datasets"],
            ["A field supervisor", "A shared dashboard link, on a phone or a laptop"],
            ["An IT administrator", "Administration, connections, projects"],
        ],
    )
    page_break(document)


def chapter_signing_in(document: Document) -> None:
    document.add_heading("2. Signing in and finding your way around", level=1)
    para(
        document,
        "Open the address your administrator gave you and sign in with your "
        "email address and password. If this is your first time, you will be "
        "asked to choose a new password before you go any further.",
    )
    figure(document, "01-sign-in", "The sign-in page.")

    document.add_heading("The parts of the screen", level=2)
    para(
        document,
        "Everything is reached from the dark bar down the left. It does not "
        "change, so you can always get back.",
    )
    table(
        document,
        ["Where", "What is there"],
        [
            [
                "**Overview**",
                "The state of everything at a glance: recent imports, indicators, failing checks",
            ],
            ["**Projects**", "Groups of work, and who may see each one"],
            ["**Datasets**", "Every file that has been imported, and the boundary layers"],
            ["**Connections**", "Survey Solutions servers and their import schedules"],
            ["**Explore**", "Building tabulations and charts"],
            ["**Dashboards**", "Assembling and reading dashboards"],
            ["**Monitoring**", "Indicators and their targets"],
            ["**Data quality**", "Checks and their results"],
            ["**Alerts**", "What has crossed a threshold, and what to do about it"],
            ["**Administration**", "Users, roles and system settings"],
        ],
    )
    figure(document, "02-overview", "The Overview, which is where signing in leaves you.")
    page_break(document)


def chapter_projects(document: Document) -> None:
    document.add_heading("3. Projects: who can see what", level=1)
    para(
        document,
        "A project is a group of work: one survey, one round, one bureau. "
        "Datasets, dashboards, indicators and boundary layers all belong to a "
        "project, or to the **shared area**, which everybody can see.",
    )
    para(
        document,
        "The shared area is a real place rather than the absence of one. A "
        "national frame of enumeration areas usually belongs there, because "
        "every survey works to it; the interviews from one round belong to that "
        "round's project.",
    )
    figure(document, "03-projects", "Projects, each showing what it holds and who is on it.")

    document.add_heading("Giving someone one project only", level=2)
    numbered(
        document,
        [
            "Open **Administration → Users** and create the account, or find it.",
            "Set the account to **Restricted**, which means it sees nothing except "
            "the projects it is added to.",
            "Open the project, go to **Members**, and add them with a role.",
        ],
    )
    note(
        document,
        "A restricted user who asks for something outside their projects is told "
        "it does not exist rather than that they may not see it. That is "
        'deliberate: a "forbidden" answer would confirm that the thing is there.',
    )

    document.add_heading("Deleting a project", level=2)
    para(
        document,
        "Deleting a project releases its datasets back to the shared area rather "
        "than destroying them. If you really do want the data gone as well, the "
        "confirmation offers that as a separate choice.",
    )
    page_break(document)


def chapter_data_in(document: Document) -> None:
    document.add_heading("4. Getting data in", level=1)

    document.add_heading("Uploading a file", level=2)
    para(
        document,
        "**Datasets → Upload data** takes Stata (`.dta`), SPSS (`.sav`), CSV, "
        "tab-separated, Excel, and zip archives of any of those. Choose the "
        "project it belongs to, give it a description if it needs one, and the "
        "file is read straight away.",
    )
    para(
        document,
        "A Stata or SPSS file brings its variable labels and value labels with "
        "it, so a column of codes shows as **North** and **South** rather than "
        "as 1 and 2 everywhere in the platform.",
    )
    figure(document, "04-datasets", "The Datasets page, grouped by project.")

    document.add_heading("An archive with several files in it", level=2)
    para(
        document,
        "A Survey Solutions export holds one file per roster level: the "
        "interview, the household members, the people abroad. Those are "
        "different tables, not different rounds, so uploading the archive makes "
        "one dataset per file inside it.",
    )
    para(
        document,
        "Uploading a later archive sends each of its files to the dataset "
        "already holding that file name. **Replace** (the default) swaps that "
        "dataset's data while keeping its identity, so the charts, indicators "
        "and quality rules built on it keep working. **Append** adds the new "
        "rows to what is there.",
    )

    document.add_heading("A questionnaire that changed mid-fieldwork", level=2)
    para(
        document,
        "When a questionnaire is revised in the field, Survey Solutions gives "
        "each version its own export. Upload them together: the first version "
        "is handled with the mode you chose and the rest are appended onto it, "
        "so all versions end up in one dataset. Name the **version column** and "
        "each row is stamped with the version it came from, which is what lets "
        "you tabulate by it afterwards.",
    )
    note(
        document,
        "This is the one place where getting the mode wrong is expensive: with "
        "Replace applied to every version in turn, version 2 would wipe version "
        "1 and version 3 would wipe version 2. Only the first version honours "
        "the mode for exactly that reason.",
    )

    document.add_heading("Importing from Survey Solutions", level=2)
    para(
        document,
        "**Connections → New connection** takes the server address, a workspace, "
        "and the credentials of an API user. Use an API user rather than an "
        "administrator: it can export data and nothing else.",
    )
    numbered(
        document,
        [
            "Enter the server, workspace, user and password, and test the connection.",
            "Choose the questionnaire and which interview statuses to bring in.",
            "Set a schedule, or leave it manual and press Import when you want it.",
        ],
    )
    figure(document, "12-connections", "Connections to Survey Solutions servers.")
    page_break(document)


def chapter_dataset(document: Document) -> None:
    document.add_heading("5. Looking at a dataset", level=1)
    para(
        document,
        "Clicking a dataset opens it: how many rows and columns, when it was "
        "last imported, and every variable with its type, its label and how "
        "much of it is missing.",
    )
    figure(document, "04b-dataset-detail", "A dataset and its variables.")

    document.add_heading("Naming variables and their codes", level=2)
    para(
        document,
        "A variable called `hh_prov_cd` means nothing to a minister. Give it a "
        "label here and that label is what appears on every axis, legend and "
        "filter from then on. The same goes for value labels: 1 and 2 become "
        "Male and Female wherever they are shown.",
    )

    document.add_heading("The command box", level=2)
    para(
        document,
        "For the changes that are quicker typed than clicked, the command box "
        "takes Stata-like commands against the dataset: `gen`, `egen`, "
        "`replace`, `label`, `drop`, `rename`, `recode`. They are checked before "
        "they run, and what they change is recorded, so a derived variable can "
        "be traced back to the line that made it.",
    )
    page_break(document)


def chapter_explore(document: Document) -> None:
    document.add_heading("6. Explore: tabulations and charts", level=1)
    para(
        document,
        "**Explore** is where a question is turned into a table or a chart. "
        "Pick a dataset, choose what to group by and what to measure, and run "
        "it. Nothing here is saved until you save it, so it is a safe place to "
        "try things.",
    )

    document.add_heading("Tabulate and chart", level=2)
    table(
        document,
        ["Box", "What it decides"],
        [
            ["**Group by**", "The variable or variables the rows are broken down by"],
            ["**Measure**", "Count, share of total, sum, mean, median, minimum or maximum"],
            ["**Filters**", "Which rows are counted at all"],
            ["**Display**", "Chart type, ordering, top-N, labels, a target line"],
        ],
    )
    para(
        document,
        "**Suggested analyses** across the top are built from the dataset's own "
        "variables, so they are a quick way into a file you have not seen before.",
    )
    figure(document, "05-explore", "Explore, with a distribution of interviews by province.")

    document.add_heading("Cross-tabulation", level=2)
    para(
        document,
        "The second tab crosses two variables. Choose a row variable, a column "
        "variable, or just one of them: a table of one variable's frequencies is "
        "a perfectly ordinary thing to want, and the platform does not make you "
        "invent a second variable to get it.",
    )
    para(
        document,
        "Percentages can be of the row, the column or the total, and a chi-square "
        "test of independence is reported when there are two variables to test.",
    )
    figure(document, "05b-crosstab", "A cross-tabulation with column percentages.")

    document.add_heading("Saving what you built", level=2)
    para(
        document,
        "**Save** keeps the query, not its results. When the dashboard shows it "
        "tomorrow it re-runs against whatever has been imported by then, which "
        "is the whole point of a monitoring dashboard.",
    )
    page_break(document)


def chapter_dashboards(document: Document) -> None:
    document.add_heading("7. Dashboards", level=1)
    para(
        document,
        "A dashboard is a page of widgets that re-runs itself. **Dashboards → "
        "New dashboard**, give it a name and a project, and it is ready for its "
        "first widget.",
    )
    figure(document, "06-dashboards", "The list of dashboards.")
    figure(document, "07-dashboard", "A fieldwork dashboard: charts, filters, pages and a map.")

    document.add_heading("The widgets", level=2)
    table(
        document,
        ["Widget", "Shows"],
        [
            [
                "**Saved chart or cross-tab**",
                "Something saved from Explore, re-run against current data",
            ],
            ["**Indicator tile**", "One tracked number with its target, status colour and trend"],
            ["**Data quality panel**", "The last result of every check on a dataset"],
            ["**Map**", "GPS points from a dataset, on a base map, with boundaries under them"],
            ["**Text note**", "A heading, an explanation, a caveat"],
            ["**Countdown**", "Time remaining to a deadline"],
            ["**Embedded HTML**", "Whatever you paste, in a sandboxed frame"],
            [
                "**How recent the data is**",
                "When each dataset was imported, and how old its newest record is",
            ],
        ],
    )
    figure(document, "15-add-widget", "Adding a widget.")

    document.add_heading("Filters", level=2)
    para(
        document,
        "**Filters** puts dropdowns above the board. A filter names a variable "
        "rather than a dataset, so it narrows every widget whose data carries "
        "that variable and leaves the rest alone.",
    )
    para(
        document,
        "Give each one a **label**. A board built for a minister should say "
        "Province, not `hh_prov_cd`.",
    )
    figure(document, "13-filters-dialog", "Choosing which filters a page offers.")

    document.add_heading("Filtering by clicking", level=2)
    para(
        document,
        "Clicking a bar, a slice, a point or a table row filters the rest of the "
        "page by what it stands for. The widget you clicked is left unfiltered, "
        "because narrowing it to the one bar you just chose would take away the "
        "means of choosing another. Click the same mark again, or the **Clear** "
        "button, to undo it.",
    )

    document.add_heading("Pages", level=2)
    para(
        document,
        "A dashboard can have several pages, each with its own widgets and its "
        "own filters, because different pages ask different questions. Use "
        "**+ Page** to add one and **Rename** to name it; the arrows move the "
        "current page left or right.",
    )
    figure(
        document,
        "08-dashboard-quality-page",
        "A second page: a cross-tab, indicator tiles and the quality panel.",
    )

    document.add_heading("Making it yours", level=2)
    para(
        document,
        "**Appearance** dresses the whole dashboard: a background colour or "
        "image, the canvas width, how fine the grid is, how transparent the "
        "widgets are, and the colour behind the page tabs.",
    )
    figure(document, "14-appearance-dialog", "Dashboard appearance.")

    document.add_heading("One widget at a time", level=2)
    para(
        document,
        "The pencil on a widget sets what belongs to that widget alone: its own "
        "background colour and transparency, a font, a text colour, a face and "
        "position for its title, a shadow, and for a chart the colour it leads "
        "with. Each falls back to the dashboard when left alone, so a single "
        "tile can be lifted off a busy background without lifting all of them.",
    )
    bullets(
        document,
        [
            "**Text colour** covers the title, the caption and the labels on the "
            "chart. A value printed on top of a bar stays white, because it sits "
            "on a filled shape rather than on the page.",
            "**Title font, size and position** dress the heading on its own. A "
            "centred title suits a tile read on its own; left is easier to scan "
            "down a column of widgets.",
            "**Shadow** lifts the widget off the background, which matters on a "
            "coloured or photographic dashboard.",
            "**Caption** puts a sentence under the widget, where a figure caption "
            "goes in a report. Use it for what the reader cannot see: which rows "
            "are counted, what was excluded, where a target came from.",
        ],
    )
    figure(document, "17-widget-settings", "One widget's own settings.")
    page_break(document)


def chapter_maps(document: Document) -> None:
    document.add_heading("8. Maps and boundaries", level=1)
    para(
        document,
        "A map widget draws the GPS points from a dataset. Tell it which "
        "variables hold the latitude and the longitude, and which columns to "
        "show when a pin is clicked. Points at the same coordinate are grouped "
        "into one pin carrying a number, so a household visited three times is "
        "one pin rather than three on top of each other.",
    )
    figure(document, "07b-map-widget", "Interview locations, with the boundaries under them.")

    document.add_heading("The ground under the map", level=2)
    bullets(
        document,
        [
            "**Streets** - OpenStreetMap, the default.",
            "**Satellite** - aerial imagery, with place names drawn over it so a "
            "cluster of pins can still be read as a village.",
            "**Terrain** - contours and relief.",
        ],
    )
    para(
        document,
        "A reader can switch between them from the control in the map's top "
        "right corner, without being able to edit the widget. **Point shape** "
        "draws the pins as circles, squares, triangles, diamonds or pentagons; "
        "a shape is told apart in a photocopy, and two shades of one colour are "
        "not. The **⤢** button fills the window with the map.",
    )
    figure(document, "20-fullscreen-map", "The same map filling the window.")

    document.add_heading("Boundary layers", level=2)
    para(
        document,
        "**Datasets → Boundaries → Add boundaries** takes the frame your "
        "fieldwork is organised into - enumeration areas, districts, villages - "
        "as GeoJSON, a GeoPackage (`.gpkg`), or a shapefile zipped together with "
        "its `.dbf` and `.shx`. Whatever your GIS office exports is read as it "
        "is; nothing has to be converted first.",
    )
    para(
        document,
        "A layer left in the shared area is available to every project, which is "
        "what a national frame usually wants.",
    )
    figure(
        document,
        "18-boundaries",
        "Boundary layers, with the attributes each one carries.",
        width=5.6,
    )

    document.add_heading("Checking that a record was collected where it says it was", level=2)
    para(
        document,
        "This is what the boundaries are really for. A household listing records "
        "the enumeration area the interviewer says they were in; the device "
        "records where they actually were. Those disagreeing is one of three "
        "things, and each is worth knowing during fieldwork rather than after "
        "it: an area code typed wrong, an interviewer working the wrong area, or "
        "a boundary the field reads differently from the office.",
    )
    para(
        document,
        "On the map widget, set **Recorded area** to the variable holding the "
        "code the interviewer recorded, and **Matched against** to the attribute "
        "on the boundary layer holding the same code. Every pin is then coloured "
        "by its verdict and the legend counts them.",
    )
    table(
        document,
        ["Colour", "Meaning"],
        [
            ["Red", "The recorded area does not match the area the GPS falls in"],
            ["Amber", "The point is outside every area in the layer"],
            ["Purple", "No area was recorded on this record"],
            ["Green", "The recorded area agrees with the GPS"],
        ],
    )
    note(
        document,
        "Codes are compared by what identifies them rather than how they were "
        "typed: `07`, `7` and `007` are one enumeration area. A code crosses "
        "from a questionnaire to a GIS file as text on one side and a number on "
        "the other, and treating those as different would report every record in "
        "the country as a mismatch.",
    )
    page_break(document)


def chapter_monitoring(document: Document) -> None:
    document.add_heading("9. Monitoring: indicators and targets", level=1)
    para(
        document,
        "An indicator is one tracked number. **Monitoring → New indicator**: "
        "name it, pick a dataset, choose the measure, and optionally give it a "
        "target and the thresholds at which it should start worrying you.",
    )
    figure(document, "09-monitoring", "Indicators with their current values and status.")

    table(
        document,
        ["Setting", "What it does"],
        [
            [
                "**Target**",
                "Draws a progress bar and gives the number something to be measured against",
            ],
            ["**Warning threshold**", "The value at which the indicator turns amber"],
            ["**Critical threshold**", "The value at which it turns red"],
            [
                "**Direction**",
                "Whether higher is better or lower is better, which reverses the logic",
            ],
            [
                "**Breakdown variable**",
                "Lets the indicator be expanded per province, team or interviewer",
            ],
            ["**Target per group**", "A separate target for each value of the breakdown"],
            ["**Percentage**", "Makes the number a share rather than a count"],
        ],
    )
    para(
        document,
        "The datasets offered are the ones in whatever the project filter above "
        "the page is set to, because an indicator belongs to whatever project "
        "its dataset belongs to. Building one from another project's data is how "
        "an indicator quietly ends up somewhere its team will not find it.",
    )

    document.add_heading("Targets per group", level=2)
    para(
        document,
        "Once a breakdown variable is chosen, each of its values can be given "
        "its own target: 140 interviews in Shefa, 50 in Torba. On the dashboard, "
        "an indicator tile with **Show the breakdown** turned on draws each "
        "group's achievement against its own target, so a province that is "
        "behind is visible without arithmetic.",
    )
    note(
        document,
        "Clicking a bar in an indicator's breakdown filters the rest of the "
        "dashboard by that group, the same as clicking any other chart.",
    )
    page_break(document)


def chapter_quality(document: Document) -> None:
    document.add_heading("10. Data quality", level=1)
    para(
        document,
        "A quality rule is a question asked of a dataset on a schedule. "
        "**Data quality → New rule**, choose the dataset and the kind of check, "
        "and set what counts as a failure.",
    )
    table(
        document,
        ["Check", "Asks"],
        [
            ["**Missing rate**", "How much of a variable is blank"],
            ["**Value range**", "Whether values fall between a minimum and a maximum"],
            ["**Duplicates**", "Whether a key, or a combination of columns, repeats"],
            ["**Outliers**", "Whether values sit far outside the usual spread"],
            ["**Consistency**", "Whether one variable agrees with another"],
            ["**Interview duration**", "Whether interviews were too short to be real"],
            ["**GPS missing**", "Whether a location was recorded"],
            [
                "**Constant value**",
                "Whether a variable never changes, which usually means it was never asked",
            ],
        ],
    )
    figure(document, "10-quality", "Quality rules and their most recent results.")
    para(
        document,
        "Failing checks are listed first and in full; passing ones are counted. "
        "A panel that lists thirty green rows buries the one red row, which is "
        "the only one anybody needs.",
    )
    para(
        document,
        "The same picture goes on a dashboard as a **data quality panel**, so a "
        "board watched by a survey manager shows whether the data behind it can "
        "be trusted.",
    )
    page_break(document)


def chapter_alerts(document: Document) -> None:
    document.add_heading("11. Alerts", level=1)
    para(
        document,
        "An alert rule watches an indicator or a quality check and raises "
        "something when it crosses a line. Alerts collect on the **Alerts** "
        "page, where each can be acknowledged so the rest of the team knows it "
        "is being dealt with, and resolved when it is.",
    )
    figure(document, "11-alerts", "The alerts page.")
    para(
        document,
        "Where email has been configured by an administrator, an alert can also "
        "be sent to a list of addresses when it is raised.",
    )
    page_break(document)


def chapter_sharing(document: Document) -> None:
    document.add_heading("12. Sharing a dashboard", level=1)
    para(
        document,
        "**Share** lists every address a dashboard is published at, and adds "
        "more. One board usually goes to several audiences at once - a minister, "
        "the field supervisors, a donor - and those do not end together, so each "
        "gets its own link.",
    )
    figure(document, "16-share-links", "Several links to one dashboard.")

    table(
        document,
        ["Action", "What it does"],
        [
            ["**Copy**", "Puts the address on the clipboard"],
            ["**Rename**", "Names the link so you can tell your audiences apart"],
            ["**Set password**", "Asks readers for a password before showing anything"],
            ["**Close**", "Stops the link working, keeping the address so it can be reopened"],
            ["**Delete**", "Destroys the link permanently"],
        ],
    )
    note(
        document,
        "Close a link rather than deleting it. A closed link can be switched "
        "back on at the same address, which matters once it is already pasted "
        "into somebody's email. Each link also records how often it has been "
        "opened, so one nobody uses can be recognised.",
    )

    document.add_heading("Passwords on a link", level=2)
    para(
        document,
        "A reader of a protected link is asked for the password once and then "
        "reads the dashboard normally. It is remembered for that browser tab "
        "only, so a shared computer in a field office does not leave the next "
        "person holding it.",
    )
    para(
        document,
        "This is not an account. Everyone holding the password is the same "
        "anonymous reader, and it exists so that a forwarded link is not a "
        "public one.",
    )
    figure(
        document,
        "22-password-door",
        "What a reader sees before giving the password.",
        width=5.4,
    )

    document.add_heading("What a reader sees", level=2)
    para(
        document,
        "A shared dashboard is the survey team's page, not the platform's: it "
        "opens with your logo and your title, with no sign-in bar above it. "
        "Readers can use the filters, click charts to narrow the page, switch "
        "the map's base layer, and copy tables and charts, but they can change "
        "nothing.",
    )
    figure(
        document,
        "21-shared-dashboard",
        "The same dashboard as a shared link, seen by somebody with no account.",
    )

    document.add_heading("Giving a dashboard its own web address", level=2)
    para(
        document,
        "A dashboard can answer on a hostname of its own, such as "
        "`labour-force.statistics.gov.vu`. Your administrator points the name at "
        "the server; you then set it on the dashboard, and the page answers "
        "there with nothing of the platform around it.",
    )
    page_break(document)


def chapter_taking_away(document: Document) -> None:
    document.add_heading("13. Taking your work away", level=1)

    document.add_heading("Copying a widget", level=2)
    para(
        document,
        "Hovering a widget shows a **⧉** button, for readers as well as authors.",
    )
    bullets(
        document,
        [
            "On a **table** it copies the rows: paste into Excel and they land in "
            "columns, paste into a plain editor and they arrive tab-separated.",
            "On a **chart** it copies a picture, ready to paste into a report. "
            "Where a browser refuses to put an image on the clipboard, the "
            "picture is saved as a file instead.",
        ],
    )
    para(
        document,
        "What is copied is what is on screen, so a table narrowed by its column "
        "filters copies narrowed.",
    )

    document.add_heading("Filtering a table by its columns", level=2)
    para(
        document,
        "The **⌕** button above a table opens a box under each heading. Typing "
        "in one narrows the rows on screen, and the button then reads how many "
        "of how many are showing. Text matches what the cell reads as, so a "
        "label matches rather than the code behind it; a number column also "
        "takes a comparison such as `> 100`.",
    )
    figure(document, "19-crosstab-widget", "A cross-tabulation on a dashboard.", width=3.9)

    document.add_heading("Downloading a dataset", level=2)
    para(
        document,
        "A dataset can be taken out whole as Stata, CSV or Excel. Stata is worth "
        "choosing when the labels matter, because they ride along with the file.",
    )
    page_break(document)


def chapter_admin(document: Document) -> None:
    document.add_heading("14. Administration", level=1)
    para(
        document,
        "**Administration** is where accounts are created, roles are set, and "
        "system settings are changed. It is visible only to administrators.",
    )
    document.add_heading("Roles", level=2)
    table(
        document,
        ["Role", "May"],
        [
            ["**Viewer**", "Read dashboards and datasets"],
            ["**Analyst**", "Also build charts, dashboards and indicators"],
            ["**Manager**", "Also import data, manage connections, quality rules and boundaries"],
            ["**Admin**", "Everything, including accounts and system settings"],
        ],
    )
    para(
        document,
        "A role is what somebody may do. **Restricted** is a separate switch "
        "saying what they may see: a restricted account sees only the projects "
        "it has been added to.",
    )
    document.add_heading("API keys", level=2)
    para(
        document,
        "For a script that needs to talk to the platform, an API key is safer "
        "than an account password: it carries a role of its own, it can be "
        "revoked on its own, and it is shown once when it is made.",
    )
    page_break(document)


def chapter_reference(document: Document) -> None:
    document.add_heading("15. Reference tables", level=1)

    document.add_heading("Chart types", level=2)
    table(
        document,
        ["Type", "Best for"],
        [
            ["Bar, horizontal bar", "Comparing categories; horizontal when the labels are long"],
            ["Stacked bar", "Parts of a whole across categories"],
            ["Population pyramid", "Two groups facing each other by age"],
            ["Line, area", "A quantity over time"],
            ["Pie, donut", "Parts of one whole, for a handful of categories"],
            ["Scatter", "Two measures against each other"],
            ["Heatmap", "Two categories crossed, read by colour"],
            ["Funnel", "Stages that narrow"],
            ["Table", "When the numbers themselves are the point"],
        ],
    )
    note(
        document,
        "Charts never use two different value axes. A second axis makes the "
        "crossing point of two lines look like a fact when it is an artefact of "
        "the scales chosen, so a second measure gets a second chart instead.",
    )

    document.add_heading("Measures", level=2)
    table(
        document,
        ["Measure", "Meaning"],
        [
            ["Count", "How many rows"],
            ["Share of total (%)", "That count as a percentage of all rows"],
            ["Sum", "The total of a numeric variable"],
            ["Mean, median", "The average, and the middle value"],
            ["Minimum, maximum", "The smallest and largest value"],
        ],
    )

    document.add_heading("Boundary file formats", level=2)
    table(
        document,
        ["Format", "What to upload"],
        [
            ["GeoJSON", "A single `.geojson` or `.json` file"],
            ["GeoPackage", "A single `.gpkg` file"],
            ["Shapefile", "A `.zip` holding the `.shp`, `.dbf` and `.shx` together"],
        ],
    )
    note(
        document,
        "Reading a boundary file needs no GIS software on the server and no "
        "internet at the moment it is read. That is deliberate: this platform is "
        "meant to run in a statistics office behind a ministry firewall.",
    )
    page_break(document)


def chapter_trouble(document: Document) -> None:
    document.add_heading("16. If something goes wrong", level=1)
    table(
        document,
        ["What you see", "What it usually is"],
        [
            [
                "A widget says a filter was ignored",
                "That widget's dataset has no such variable. The filter applied to "
                "the widgets that do have it and left this one alone, which is "
                "what it is meant to do.",
            ],
            [
                "The map draws pins but no map",
                "The server cannot reach the tile host. The pins are the data and "
                "they still draw. A deployment with its own tile service can be "
                "pointed at it with **Map tiles**.",
            ],
            [
                "Every record reports as a mismatch",
                "The **Matched against** attribute is probably the wrong column on "
                "the boundary layer. Check which of its attributes actually holds "
                "the area code.",
            ],
            [
                "An import failed",
                "Open the dataset: the reason is recorded against it. A file whose "
                "columns changed shape is the usual cause.",
            ],
            [
                "A dashboard is slow",
                "Usually one widget over a very large dataset. Narrow it with a "
                "filter, or reduce how many categories it draws with top-N.",
            ],
            [
                "A shared link says it is not available",
                "It has been closed, or deleted. Closed links can be reopened from "
                "**Share** on the dashboard.",
            ],
        ],
    )

    document.add_heading("Getting help", level=2)
    para(
        document,
        "The full technical documentation - installation, backups, upgrades and "
        "the API - lives with the source code in the `docs` folder: "
        "`deployment.md`, `architecture.md`, `api.md` and `survey-solutions.md`. "
        "This manual covers using the platform; those cover running it.",
    )


def build() -> Path:
    document = Document()
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
    style(document)
    footer(document)

    cover(document)
    contents(document)
    chapter_what(document)
    chapter_signing_in(document)
    chapter_projects(document)
    chapter_data_in(document)
    chapter_dataset(document)
    chapter_explore(document)
    chapter_dashboards(document)
    chapter_maps(document)
    chapter_monitoring(document)
    chapter_quality(document)
    chapter_alerts(document)
    chapter_sharing(document)
    chapter_taking_away(document)
    chapter_admin(document)
    chapter_reference(document)
    chapter_trouble(document)

    document.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size / 1024 / 1024:.1f} MB, {len(figures)} figures)")
